#!/usr/bin/env python3
"""아이템 목록: 엑셀(.xlsx) -> Godot용 JSON 변환기.

data/items.xlsx 를 읽어 data/items.json (UTF-8) 을 생성합니다.

    사용법:  python tools/convert_items.py

필요 패키지:  pip install openpyxl

엑셀 첫 행은 헤더(열 이름). 권장 열(대소문자 무시):
    id            : 아이템 식별자(영문) — 필수 (예: potato)
    display_name  : 표시 이름(한글 가능)
    category      : food | material | relic | clue
    max_stack     : 한 칸에 겹칠 수 있는 최대 개수(기본 99)
    note          : 설명(게임 미사용)
"""
import json
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("[에러] openpyxl 이 필요합니다.  pip install openpyxl")
    sys.exit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "data", "items.xlsx")
OUT = os.path.join(ROOT, "data", "items.json")

NUMERIC = {"max_stack"}


def coerce_num(v):
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip()
    try:
        return int(s)
    except ValueError:
        return s


def main() -> int:
    if not os.path.exists(XLSX):
        print(f"[에러] 엑셀 파일 없음: {XLSX}")
        return 1

    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("[에러] 시트가 비어 있습니다.")
        return 1

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    if "id" not in header:
        print(f"[에러] 헤더에 'id' 열이 필요합니다. 현재 헤더: {header}")
        return 1

    items = []
    for r in rows[1:]:
        if r is None:
            continue
        rec = {}
        for i, col in enumerate(header):
            if col == "" or col == "note":
                continue
            val = r[i] if i < len(r) else None
            if val is None:
                continue
            rec[col] = coerce_num(val) if col in NUMERIC else str(val).strip()
        if not rec.get("id"):
            continue
        items.append(rec)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    print(f"[완료] 아이템 {len(items)}종 → {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
