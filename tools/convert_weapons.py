#!/usr/bin/env python3
"""무기 카탈로그: 엑셀(.xlsx) -> Godot용 JSON 변환기.

data/weapons.xlsx 를 읽어 data/weapons.json (UTF-8) 을 생성합니다.

    사용법:  python tools/convert_weapons.py

필요 패키지:  pip install openpyxl

엑셀 첫 행은 헤더(열 이름). 권장 열(대소문자 무시):
    id            : 무기 식별자(영문) — 필수 (예: pistol, rifle)
    display_name  : 표시 이름
    kind          : ranged | melee
    base_damage   : 기본 공격력(강화 레벨 보너스는 별도)
    craft_cost    : 제작 비용(토큰). 0 이면 기본 보유 무기
    uses_ammo     : 탄약 사용 여부(TRUE/FALSE)
    note          : 설명(게임 미사용)

행 순서 = 표시 순서. craft_cost=0 인 무기는 새 게임 시 기본 보유.
"""
import json
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("[에러] openpyxl 이 필요합니다.  pip install openpyxl")
    sys.exit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "data", "weapons.xlsx")
OUT = os.path.join(ROOT, "data", "weapons.json")

NUMERIC = {"base_damage", "craft_cost", "muzzle_x"}
BOOLEAN = {"uses_ammo"}


def coerce_num(v):
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip()
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return s


def coerce_bool(v):
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in ("1", "true", "yes", "y", "o")


def main() -> int:
    if not os.path.exists(XLSX):
        print(f"[에러] 엑셀 파일 없음: {XLSX}")
        return 1

    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("[에러] 시트가 비어 있습니다.")
        return 1

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    if "id" not in header:
        print(f"[에러] 헤더에 'id' 열이 필요합니다. 현재 헤더: {header}")
        return 1

    weapons = []
    for r in rows[1:]:
        if r is None:
            continue
        rec = {}
        for i, col in enumerate(header):
            if col == "" or col == "note":
                continue
            val = r[i] if i < len(r) else None
            if val is None:
                continue
            if col in NUMERIC:
                rec[col] = coerce_num(val)
            elif col in BOOLEAN:
                rec[col] = coerce_bool(val)
            else:
                rec[col] = str(val).strip()
        if not rec.get("id"):
            continue
        weapons.append(rec)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(weapons, f, ensure_ascii=False, indent=2)
    print(f"[완료] 무기 {len(weapons)}종 → {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
