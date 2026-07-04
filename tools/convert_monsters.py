#!/usr/bin/env python3
"""몬스터 데이터: 엑셀(.xlsx) -> Godot용 JSON 변환기.

data/monsters.xlsx 를 읽어 data/monsters.json (UTF-8) 을 생성합니다.
엑셀에서 몬스터 종류/능력치를 수정한 뒤 이 스크립트를 실행하면 게임 값이 갱신됩니다.

    사용법:  python tools/convert_monsters.py

필요 패키지:  pip install openpyxl

엑셀 시트 첫 행은 헤더(열 이름)여야 합니다. 권장 열(대소문자 무시):
    id                : 몬스터 식별자(영문) — 필수 (예: monster_a)
    display_name      : 표시 이름
    max_hp            : 최대 체력
    contact_damage    : 접촉 공격력
    speed             : 추적 속도(px/s)
    detection_range   : 감지 반경(px)
    chase_duration    : 무피해 시 추적 유지 시간(초)
    wander_radius     : 대기 배회 반경(px)
    wander_speed      : 배회 속도(px/s)
    sprite            : idle 애니 폴더 id(assets/monsters/<sprite>/<sprite>_idle_N.png). 비우면 사각형 플레이스홀더
    note              : 설명(게임 미사용)

행 순서가 곧 몬스터 종류 순서입니다(스폰 시 순환). 첫 행이 '가장 첫번째 일반 몬스터'.
"""
import json
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("[에러] openpyxl 이 필요합니다.  pip install openpyxl")
    sys.exit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "data", "monsters.xlsx")
OUT = os.path.join(ROOT, "data", "monsters.json")

# 숫자로 해석할 열(그 외는 문자열 유지).
NUMERIC = {
    "max_hp", "contact_damage", "speed", "detection_range",
    "chase_duration", "wander_radius", "wander_speed",
    "stop_distance", "attack_range", "attack_cooldown", "attack_windup",
    "drop_chance", "drop_min", "drop_max", "scale",
}
BOOLEAN = {"is_boss"}


def coerce_num(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip()
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return s


def main() -> int:
    if not os.path.exists(XLSX):
        print(f"[에러] 엑셀 파일 없음: {XLSX}")
        return 1

    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("[에러] 시트가 비어 있습니다.")
        return 1

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    if "id" not in header:
        print(f"[에러] 헤더에 'id' 열이 필요합니다. 현재 헤더: {header}")
        return 1

    monsters = []
    for r in rows[1:]:
        if r is None:
            continue
        rec = {}
        for i, col in enumerate(header):
            if col == "" or col == "note":
                continue
            val = r[i] if i < len(r) else None
            if val is None:
                continue
            if col in NUMERIC:
                rec[col] = coerce_num(val)
            elif col in BOOLEAN:
                rec[col] = str(val).strip().lower() in ("1", "true", "yes", "y", "o")
            else:
                rec[col] = str(val).strip()
        if not rec.get("id"):
            continue
        monsters.append(rec)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(monsters, f, ensure_ascii=False, indent=2)
    print(f"[완료] 몬스터 {len(monsters)}종 → {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
