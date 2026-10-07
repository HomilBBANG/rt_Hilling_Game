#!/usr/bin/env python3
"""쿠킹 밸런스: 엑셀(.xlsx) -> Godot용 JSON 변환기.

data/cooking.xlsx 를 읽어 data/cooking.json (UTF-8) 을 생성합니다.

    사용법:  python tools/convert_cooking.py

필요 패키지:  pip install openpyxl

시트 구성
---------
settings    : key | value | note        — 쿠킹 전역 수치(밤 시간, 목표 만족도, 등급 기준, 토큰, 숙련도 등)
ingredients : 조리대에서 꺼낼 수 있는 식재료
    id          : 아이템 id (items.xlsx 와 같음)
    chop_count  : 조리대에서 클릭해 써는 횟수. 0 이면 손질 없이 바로 사용
    color       : 이미지가 없을 때 임시 도형 색(#RRGGBB)
recipes     : 레시피 × 레벨(1~3) 한 행씩
    id               : 레시피 식별자(영문) — 필수
    level            : 1, 2, 3 — 필수
    display_name     : 표시 이름 (레벨 1 행에만 적어도 됨)
    unlock           : 해금 조건 (레벨 1 행) — default | item:<아이템id>
    ingredients      : 필요 재료(순서 무관). 쉼표 구분, 중복 = 여러 개 필요 (예: potato, potato, herb)
    station          : 재료를 넣어 완성하는 기구 — fryer(튀김기) | pot(냄비) | bowl(믹싱볼, 불 없이 바로 완성)
    timer_seconds    : 튀김기/냄비 적정 시간(초). bowl 은 무시
    sat_a/sat_b/sat_c: 등급(A/B/C)별 만족도 증가치
    upgrade_mastery  : 다음 레벨로 강화하는 데 필요한 누적 숙련도. 마지막 레벨은 비움
    note             : 설명(게임 미사용)

행 순서 = 표시 순서.
"""
import json
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("[에러] openpyxl 이 필요합니다.  pip install openpyxl")
    sys.exit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "data", "cooking.xlsx")
OUT = os.path.join(ROOT, "data", "cooking.json")

STATIONS = {"fryer", "pot", "bowl"}


def num(v, default=0):
    if v is None or str(v).strip() == "":
        return default
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip()
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return default


def text(v):
    return "" if v is None else str(v).strip()


def read_settings(ws):
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r or r[0] is None or text(r[0]) == "":
            continue
        out[text(r[0])] = num(r[1] if len(r) > 1 else None)
    return out


def read_recipes(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [text(c).lower() for c in rows[0]]
    for need in ("id", "level"):
        if need not in header:
            raise ValueError(f"recipes 시트 헤더에 '{need}' 열이 필요합니다. 현재: {header}")

    recipes = {}
    order = []
    for r in rows[1:]:
        if r is None:
            continue
        row = {header[i]: (r[i] if i < len(r) else None) for i in range(len(header)) if header[i]}
        rid = text(row.get("id"))
        if rid == "":
            continue
        if rid not in recipes:
            recipes[rid] = {"id": rid, "display_name": rid, "unlock": "default", "levels": []}
            order.append(rid)
        rec = recipes[rid]
        if text(row.get("display_name")):
            rec["display_name"] = text(row.get("display_name"))
        if text(row.get("unlock")):
            rec["unlock"] = text(row.get("unlock"))
        ings = [s.strip() for s in text(row.get("ingredients")).split(",") if s.strip()]
        timer = num(row.get("timer_seconds"), 0)
        station = text(row.get("station")).lower()
        if station == "":
            # 구버전 시트 호환: timer_name(튀기기/끓이기)으로 조리 기구 추정, 타이머 없으면 믹싱볼
            station = ("pot" if "끓" in text(row.get("timer_name")) else "fryer") if timer else "bowl"
        if station == "bowl":
            timer = 0
        if station and station not in STATIONS:
            raise ValueError(f"{rid} Lv{row.get('level')}: station 은 {sorted(STATIONS)} 중 하나여야 합니다 (현재: {station})")
        rec["levels"].append({
            "level": int(num(row.get("level"), 1)),
            "ingredients": ings,
            "station": station,
            "timer_seconds": timer,
            "sat": {
                "A": num(row.get("sat_a"), 0),
                "B": num(row.get("sat_b"), 0),
                "C": num(row.get("sat_c"), 0),
            },
            "upgrade_mastery": int(num(row.get("upgrade_mastery"), 0)),
        })

    out = []
    for rid in order:
        rec = recipes[rid]
        rec["levels"].sort(key=lambda lv: lv["level"])
        out.append(rec)
    return out


def read_ingredients(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [text(c).lower() for c in rows[0]]
    out = []
    for r in rows[1:]:
        row = {header[i]: (r[i] if i < len(r) else None) for i in range(len(header)) if header[i]}
        iid = text(row.get("id"))
        if iid == "":
            continue
        out.append({
            "id": iid,
            "chop_count": int(num(row.get("chop_count"), 0)),
            "color": text(row.get("color")) or "#B0B0B0",
        })
    return out


def main() -> int:
    if not os.path.exists(XLSX):
        print(f"[에러] 엑셀 파일 없음: {XLSX}")
        return 1
    wb = load_workbook(XLSX, data_only=True)
    if "settings" not in wb.sheetnames or "recipes" not in wb.sheetnames:
        print(f"[에러] 'settings', 'recipes' 시트가 필요합니다. 현재: {wb.sheetnames}")
        return 1
    try:
        data = {
            "settings": read_settings(wb["settings"]),
            "ingredients": read_ingredients(wb["ingredients"]) if "ingredients" in wb.sheetnames else [],
            "recipes": read_recipes(wb["recipes"]),
        }
    except ValueError as e:
        print(f"[에러] {e}")
        return 1

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"[완료] 쿠킹 설정 {len(data['settings'])}개, 식재료 {len(data['ingredients'])}종, 레시피 {len(data['recipes'])}종 → {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
